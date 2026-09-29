"""Read spreadsheet cells for user-reviewed column mapping; never evaluate workbook code."""

from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .merged_cells import workbook_merges


def workbook_cells(content):
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ValueError("无法读取 XLSX 工作簿，请检查文件格式或是否加密") from error
    try:
        merges = workbook_merges(content)
        return dict(
            sheets=[
                dict(name=sheet.title, rows=sheet_rows(sheet))
                | ({"merges": merges[sheet.title]} if merges[sheet.title] else {})
                for sheet in workbook.worksheets
            ]
        )
    finally:
        workbook.close()


def sheet_rows(sheet):
    rows = []
    for cells in sheet.iter_rows():
        values = [cell_text(cell) for cell in cells]
        while values and values[-1] == "":
            values.pop()
        rows.append(values)
    while rows and not rows[-1]:
        rows.pop()
    return rows


def cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    # Formula text stays visible. Numeric fields must be confirmed as literal values.
    return str(value)
