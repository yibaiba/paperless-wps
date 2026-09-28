"""File adapter. Formulas and caches are derived from the same frozen quote projection."""

from copy import copy
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.workbook.properties import CalcProperties

from .calculation import project_description
from .layout import detail_rows
from .template import AUXILIARY_TITLE, SECTIONS, TEMPLATE_SHA256

STYLE_ATTRIBUTES = ("font", "fill", "border", "alignment", "number_format", "protection")
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class ExcelRenderer:
    def __init__(self, template_path):
        self.template_path = template_path

    def quotation(self, saved):
        content = self.template_path.read_bytes()
        if sha256(content).hexdigest() != TEMPLATE_SHA256:
            raise ValueError("报价模板校验值变化，请注册新的模板映射版本")
        wb = load_workbook(BytesIO(content))
        ws = wb["Sheet1"]
        styles = {
            key: [cell_style(ws.cell(row, c)) for c in range(1, 11)]
            for key, row in (("detail", 8), ("section", 7), ("total", 35))
        }
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row >= 7:
                ws.unmerge_cells(str(merged))
        ws.delete_rows(7, 29)
        quote = saved["configuration"]["quotation"]
        output = saved["quotation_output"]
        fill_header(ws, quote)
        last, cache = fill_sections(ws, output, styles=styles)
        fill_total(
            ws,
            last,
            styles=styles,
            metadata=quote,
            business_gaps=unresolved_business(output, saved["configuration"]),
        )
        literal(ws.cell(last, 9), "报价草稿")
        literal(ws.cell(last, 10), f"保存版本 {saved['revision']}")
        total = output["total"] if output["total"] is not None else "待确认"
        cache[f"H{last}"] = Decimal(total) if total != "待确认" else total
        cache[f"E{last}"] = cache[f"H{last}"]
        ws.print_area = f"A1:J{last}"
        ws.print_title_rows = "1:6"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
        buffer = BytesIO()
        wb.save(buffer)
        return formula_caches(buffer.getvalue(), cache)

    def configuration(self, saved):
        wb = Workbook()
        ws = wb.active
        ws.title = "设备与采购清单"
        ws.append(
            [
                "名称",
                "型号",
                "规格",
                "类型",
                "部署数量",
                "已有数量",
                "采购数量",
                "供货待定",
                "单位",
                "服务系统",
                "来源",
                "备注",
            ]
        )
        for line in saved["project_output"]["lines"]:
            supply = line.get("supply")
            values = [
                line["name"],
                line["model"],
                line["specification"],
                line["kind"],
                Decimal(line["quantity"]),
                Decimal(supply["existing"]) if supply else None,
                Decimal(supply["purchase"]) if supply else None,
                Decimal(supply["unknown"]) + Decimal(supply["unassigned"])
                if supply
                else "未区分供货",
                line["unit"],
                " / ".join(sorted({c["system_name"] for c in line["consumers"]})),
                f"{line['source'].get('sheet', '')} 第{line['source'].get('row', '')}行",
                line["note"],
            ]
            append_literals(ws, values)
        issues = wb.create_sheet("检查与配套")
        issues.append(["类型", "状态", "对象", "内容"])
        for item in saved["checks"] + (saved.get("quotation_output") or {}).get("issues", []):
            append_literals(
                issues,
                [
                    item["kind"],
                    item["status"],
                    item.get("device_id") or item.get("requirement_id", ""),
                    item.get("message", ""),
                ],
            )
        for item in saved["suggestions"]:
            append_literals(
                issues,
                [
                    "配套",
                    item["status"],
                    item["id"],
                    f"需求量 {item.get('required')}；缺量 {item.get('missing')}；"
                    f"{item.get('missing_information', [])}",
                ],
            )
        versions = wb.create_sheet("版本来源")
        for key in ("project_id", "revision", "fingerprint"):
            versions.append([key, saved.get(key)])
        for key in ("knowledge_snapshot_id", "definition_snapshot_id", "calculation_version"):
            versions.append([key, saved["configuration"].get(key)])
        for sheet in wb:
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in sheet:
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column in "ABCDEFGHIJKL":
                sheet.column_dimensions[column].width = 22
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


def unresolved_business(output, configuration):
    devices = {device["id"]: device for device in configuration["devices"]}
    return any(
        not line["supply_complete"]
        or project_description(
            line, device=devices[line["device_id"]], quote=configuration["quotation"]
        )[1] is not None
        for line in output["lines"]
    )


def cell_style(cell):
    return {key: copy(getattr(cell, key)) for key in STYLE_ATTRIBUTES}


def style_row(ws, row, styles):
    for col, style in enumerate(styles, 1):
        for key, value in style.items():
            setattr(ws.cell(row, col), key, copy(value))


def literal(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def append_literals(ws, values):
    row = ws.max_row + 1
    for col, value in enumerate(values, 1):
        literal(ws.cell(row, col), value)


def fill_header(ws, quote):
    cells = {
        "C4": "customer",
        "C5": "project_name",
        "H2": "sales_contact",
        "H3": "designer_contact",
        "H4": "design_date",
        "H5": "room_description",
    }
    for address, key in cells.items():
        literal(ws[address], quote.get(key) or "")


def fill_sections(ws, output, *, styles):
    groups = {name: [] for name, _ in SECTIONS}
    for line in output["lines"]:
        if Decimal(line["purchase_quantity"]) > 0 or Decimal(line["unknown_quantity"]) > 0:
            groups.setdefault(line["section"], []).append(line)
    row, number, cache = 7, 0, {}
    for name, lines in groups.items():
        style_row(ws, row, styles["section"])
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        literal(ws.cell(row, 1), AUXILIARY_TITLE if name == "其他辅助设备" else name)
        ws.row_dimensions[row].height = 25
        row += 1
        details = detail_rows(ws, lines)
        count = max(len(details), dict(SECTIONS).get(name, 1))
        for index in range(count):
            style_row(ws, row, styles["detail"])
            ws.row_dimensions[row].height = 28
            ws.cell(row, 8, f'=IF(OR(E{row}="",G{row}=""),"",ROUND(E{row}*G{row},2))')
            cache[f"H{row}"] = ""
            if index < len(details):
                detail = details[index]
                number += not detail["continuation"]
                line = detail["line"]
                fill_line(ws, row, detail=detail, number=number)
                if detail["continuation"]:
                    row += 1
                    continue
                cache[f"H{row}"] = Decimal(line["amount"]) if line["amount"] is not None else ""
                if (
                    Decimal(line["purchase_quantity"]) == 0
                    and Decimal(line["unknown_quantity"]) > 0
                ):
                    cache[f"H{row}"] = ""
            row += 1
    return row, cache


def fill_line(ws, row, *, detail, number):
    line = detail["line"]
    pending_only = Decimal(line["purchase_quantity"]) == 0 and Decimal(line["unknown_quantity"]) > 0
    values = {
        1: number,
        5: None if pending_only else Decimal(line["purchase_quantity"]),
        7: Decimal(line["unit_price"]) if line["unit_price"] is not None else None,
    }
    for col, value in {**values, **detail["texts"]}.items():
        literal(ws.cell(row, col), value)
    if detail["continuation"]:
        for col in (1, 5, 7):
            ws.cell(row, col).value = None
    ws.row_dimensions[row].height = detail["height"]


def fill_total(ws, row, *, styles, metadata, business_gaps):
    style_row(ws, row, styles["total"])
    for start, end in ((1, 2), (3, 4), (5, 7)):
        ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    literal(ws.cell(row, 1), "报价合计")
    literal(ws.cell(row, 3), metadata["tax_terms"])
    # Only actual items have numeric sequence numbers. Section titles, spare rows
    # and description continuations do not. COUNT includes explicit zero amounts.
    incomplete = f"COUNT(A8:A{row - 1})>COUNT(H8:H{row - 1})"
    # Supply and obsolete descriptions require business review, not a price edit.
    guard = f"OR({incomplete},TRUE)" if business_gaps else incomplete
    ws.cell(row, 8, f'=IF({guard},"待确认",SUM(H8:H{row - 1}))')
    ws.cell(row, 5, f"=H{row}")
    ws.row_dimensions[row].height = 42


def formula_caches(content, values):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                for cell in root.iter(f"{{{NS}}}c"):
                    if cell.get("r") in values:
                        value = values[cell.get("r")]
                        cached = cell.find(f"{{{NS}}}v")
                        if cached is None:
                            cached = ET.SubElement(cell, f"{{{NS}}}v")
                        cell.set("t", "str" if isinstance(value, str) else "n")
                        cached.text = str(value)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, data)
    return output.getvalue()
