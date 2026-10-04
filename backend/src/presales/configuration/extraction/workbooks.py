"""Range-scoped workbook evidence; formulas remain literal source text."""

from hashlib import sha256

from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from sqlalchemy import select

from presales.lists.receipts import once
from presales.quotation.importing import workbook_cells
from presales.rules.calculation import digest
from presales.rules.repository import RuleConflict

from ..common import Entities
from ..models import Revision


def workbook_preview(content, *, filename):
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("请选择 .xlsx 工作簿")
    return dict(name=filename, digest=sha256(content).hexdigest(), **workbook_cells(content))


def cell_value(rows, row, column):
    return rows[row - 1][column - 1] if row <= len(rows) and column <= len(rows[row - 1]) else ""


def cell_segment(sheet, row, column):
    address = f"{get_column_letter(column)}{row}"
    merge = next(
        (
            m
            for m in sheet.get("merges", [])
            if m["start_row"] <= row <= m["end_row"]
            and m["start_column"] <= column <= m["end_column"]
        ),
        None,
    )
    anchor_row, anchor_column = (
        (merge["start_row"], merge["start_column"]) if merge else (row, column)
    )
    return dict(
        id=digest([sheet["name"], address]),
        sheet=sheet["name"],
        cell=address,
        anchor=f"{get_column_letter(anchor_column)}{anchor_row}",
        merge_range=merge["range"] if merge else None,
        location=f"{sheet['name']}!{address}",
        text=cell_value(sheet["rows"], anchor_row, anchor_column),
    )


def selected_segments(preview, ranges):
    sheets = {s["name"]: s for s in preview["sheets"]}
    segments = {}
    for selection in ranges:
        sheet = sheets.get(selection.sheet)
        if sheet is None:
            raise ValueError("工作表不存在：" + selection.sheet)
        region = CellRange(selection.range)
        for row, column in populated_region(sheet, region):
            segment = cell_segment(sheet, row, column)
            segments[segment["id"]] = segment
    if not segments:
        raise ValueError("选择范围内没有资料内容")
    return list(segments.values())


def populated_region(sheet, region):
    # Iterate the actual worksheet extent, not formatted or user-selected empty tails.
    merges = sheet.get("merges", [])
    last_row = max([len(sheet["rows"]), *[m["end_row"] for m in merges]])
    last_column = max([0, *map(len, sheet["rows"]), *[m["end_column"] for m in merges]])
    for row in range(region.min_row, min(region.max_row, last_row) + 1):
        for column in range(region.min_col, min(region.max_col, last_column) + 1):
            yield row, column


def save_workbook(session, data, *, preview):
    if preview["digest"] != data.digest:
        raise RuleConflict("工作簿与预览不同，请重新核对文件")
    segments = selected_segments(preview, data.ranges)
    payload = dict(
        name=data.name,
        format="xlsx",
        filename=preview["name"],
        digest=data.digest,
        ranges=[r.model_dump() for r in data.ranges],
        segments=segments,
        actor=data.actor,
        evidence=data.evidence,
    )
    return once(
        session,
        namespace="workbook_material",
        request=data,
        perform=lambda: Entities(session).save(
            "material",
            payload,
            entity_id=data.material_id,
            expected_revision=data.expected_revision,
        ),
    )


def material_revision(session, identity, revision):
    Entities(session).get(identity, kind="material")
    record = session.scalar(
        select(Revision).where(
            Revision.entity_id == identity,
            Revision.revision == revision,
        )
    )
    if record is None:
        raise ValueError("引用的资料修订不存在")
    return dict(id=identity, revision=revision, **record.payload)


def referenced_segment(session, reference):
    material = material_revision(session, reference.material_id, reference.material_revision)
    segment = next((s for s in material["segments"] if s.get("id") == reference.segment_id), None)
    if segment is None:
        raise ValueError("引用的资料片段不存在于该修订")
    return segment
