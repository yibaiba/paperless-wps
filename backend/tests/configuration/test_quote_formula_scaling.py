from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from presales.quotation.excel import ExcelRenderer
from presales.quotation.template import TEMPLATE_PATH

from .test_list_mcp import saved_project


@pytest.mark.parametrize("case", ["zero", "missing_price", "unknown_supply"])
def test_large_quote_uses_range_formula_and_keeps_unknowns(client, catalog, case):
    saved = saved_project(client, catalog)
    record = client.get("/api/configuration/projects/" + saved["project_id"]).json()
    prototype = record["quotation_output"]["lines"][0]
    lines = [dict(deepcopy(prototype), unit_price="0", amount="0.00") for _ in range(500)]
    # Continuations must not look like another chargeable item to the total formula.
    lines[0]["specification"] = "很长的产品规格说明\n" * 80
    if case == "missing_price":
        lines[-1].update(unit_price=None, amount=None)
    if case == "unknown_supply":
        lines[-1].update(supply_complete=False, unknown_quantity="1")
    total = "0.00" if case == "zero" else None
    record["quotation_output"].update(lines=lines, total=total)
    content = ExcelRenderer(TEMPLATE_PATH).quotation(record)
    formulas = load_workbook(BytesIO(content)).active
    values = load_workbook(BytesIO(content), data_only=True).active
    last = formulas.max_row
    formula = formulas.cell(last, 8).value
    assert f"COUNT(A8:A{last - 1})>COUNT(H8:H{last - 1})" in formula
    assert len(formula) < 150
    assert (",TRUE)" in formula) == (case == "unknown_supply")
    assert sum(isinstance(values.cell(r, 1).value, int) for r in range(8, last)) == 500
    assert values.cell(last, 8).value == (0 if total is not None else "待确认")
    assert len(formulas._images) == 1
