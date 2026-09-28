from copy import deepcopy
from io import BytesIO

from openpyxl import Workbook, load_workbook

from presales.quotation.excel import ExcelRenderer
from presales.quotation.template import TEMPLATE_PATH

from .test_list_mcp import call, mutation, saved_project
from .test_sheet_edits import preview


def test_import_cells_does_not_execute_formulas_or_write_catalog(client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户原表"
    sheet.append(["名称", "型号", "数量", "单价"])
    sheet.append(["多行\n描述", "SERVER-X", 2.5, "=1+2"])
    data = BytesIO()
    workbook.save(data)
    before = client.get("/api/configuration/variants").json()
    result = client.post(
        "/api/quotation/import-cells", files={"file": ("quote.xlsx", data.getvalue())}
    )
    assert result.status_code == 200, result.text
    assert result.json()["sheets"] == [
        dict(
            name="客户原表",
            rows=[
                ["名称", "型号", "数量", "单价"],
                ["多行\n描述", "SERVER-X", "2.5", "=1+2"],
            ],
        )
    ]
    assert client.get("/api/configuration/variants").json() == before
    for name in ("bad.xlsx", "old.xls"):
        assert (
            client.post(
                "/api/quotation/import-cells", files={"file": (name, b"invalid")}
            ).status_code
            == 422
        )


def test_purchase_edit_preserves_owned_and_deployment_reports_overallocation(client, catalog):
    saved = saved_project(client, catalog)
    current = client.get("/api/configuration/projects/" + saved["project_id"]).json()[
        "configuration"
    ]
    current["supply_allocations"].append(
        dict(id="owned", device_id="server", quantity="0.5", source="existing", evidence="客户已有")
    )
    operations = [
        dict(action="purchase_set", device_id="server", quantity="3.25", evidence="测试采购量")
    ]
    response = preview(client, saved, operations, configuration=current)
    assert response.status_code == 200, response.text
    checked = response.json()["checked"]
    assert checked["configuration"]["devices"][0]["quantity"] == "2.5"
    assert checked["quotation_output"]["lines"][0]["supply"]["existing"] == "0.5"
    assert checked["quotation_output"]["lines"][0]["purchase_quantity"] == "3.25"
    assert checked["quotation_output"]["total"] is None
    assert any("超过" in issue for issue in checked["quotation_output"]["lines"][0]["issues"])
    operations[0]["quantity"] = "0"
    zero = preview(client, saved, operations, configuration=current).json()["checked"]
    assert zero["quotation_output"]["lines"][0]["purchase_quantity"] == "0"
    assert zero["configuration"]["supply_allocations"] == [current["supply_allocations"][-1]]


def test_import_blank_price_never_adopts_catalog_zero_is_explicit(client, catalog):
    saved = saved_project(client, catalog)
    price = dict(
        device_id="server",
        variant_id=catalog["variants"][0]["id"],
        source_id=catalog["sources"][0]["id"],
        mode="import",
        unit_price=None,
        evidence="原表缺价",
        price_column="",
    )
    missing = preview(client, saved, [dict(action="price_set", value=price)]).json()["checked"]
    assert missing["quotation_output"]["total"] is None
    assert missing["quotation_output"]["lines"][0]["unit_price"] is None
    repeated = client.post(
        "/api/configuration/check", json=dict(configuration=missing["configuration"])
    ).json()
    assert repeated["quotation_output"]["total"] is None
    price["unit_price"] = "0"
    zero = preview(client, saved, [dict(action="price_set", value=price)]).json()["checked"]
    assert zero["quotation_output"]["total"] == "0.00"


def test_description_web_mcp_save_export_original_and_stale(client, catalog):
    saved = saved_project(client, catalog)
    project_id = saved["project_id"]
    original = client.get("/api/configuration/projects/" + project_id).json()["configuration"]
    operations = [
        dict(action="description_set", device_id="server", text="项目专用说明\n不修改产品库"),
        dict(action="purchase_set", device_id="server", quantity="1.5", evidence="测试"),
    ]
    web = preview(client, saved, operations).json()["checked"]
    line = web["quotation_output"]["lines"][0]
    assert line["specification"] == operations[0]["text"]
    assert (
        line["original_specification"] == original["devices"][0]["source_snapshot"]["specification"]
    )
    assert (
        web["configuration"]["devices"][0]["source_snapshot"]
        == original["devices"][0]["source_snapshot"]
    )
    draft = call(
        client,
        "list_create",
        dict(
            name="说明测试",
            project_id=project_id,
            revision=1,
            operation_id="description-copy",
            actor="测试",
            evidence="隔离",
        ),
    )
    draft = call(client, "list_update", mutation(draft, operations=operations))
    quote = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert quote["items"][0]["specification"] == line["specification"]
    response = client.put(
        "/api/configuration/projects/" + project_id,
        json=dict(expected_revision=1, configuration=web["configuration"]),
    )
    assert response.status_code == 200, response.text
    reopened = client.get("/api/configuration/projects/" + project_id).json()
    assert (
        reopened["configuration"]["quotation"]["descriptions"]["server"]["text"]
        == line["specification"]
    )
    exported = load_workbook(BytesIO(ExcelRenderer(TEMPLATE_PATH).quotation(reopened)))
    assert any("项目专用说明" in str(cell.value) for row in exported.active for cell in row)
    changed = deepcopy(reopened["configuration"])
    device = changed["devices"][0]
    replacement = dict(
        id=device["id"],
        name=device["name"],
        quantity=device["quantity"],
        kind=device["kind"],
        variant_id=catalog["variants"][1]["id"],
        source_id=catalog["sources"][1]["id"],
    )
    stale = preview(
        client,
        saved,
        [dict(action="device_put", value=replacement)],
        configuration=changed,
        expected_revision=2,
    ).json()["checked"]
    assert any(
        "项目产品说明已过期" in issue for issue in stale["quotation_output"]["lines"][0]["issues"]
    )
    restored = preview(
        client,
        saved,
        [dict(action="description_set", device_id="server", text=None)],
        configuration=changed,
        expected_revision=2,
    ).json()["checked"]
    assert (
        restored["quotation_output"]["lines"][0]["specification"] == line["original_specification"]
    )


def test_failed_import_batch_does_not_modify_saved_project(client, catalog):
    saved = saved_project(client, catalog)
    response = preview(
        client,
        saved,
        [
            dict(action="description_set", device_id="server", text="不会保存"),
            dict(action="purchase_set", device_id="missing", quantity="2", evidence="测试"),
        ],
    )
    assert response.status_code == 422
    assert response.json()["detail"]["operation_index"] == 1
    assert (
        client.get("/api/configuration/projects/" + saved["project_id"]).json()["configuration"][
            "quotation"
        ]["descriptions"]
        == {}
    )


def test_import_two_same_model_lines_stay_separate_and_retry_is_idempotent(client, catalog):
    from .test_list_mcp import create, device_operations

    draft = call(
        client, "list_update", mutation(create(client), operations=device_operations(catalog))
    )
    variant, source = catalog["variants"][0], catalog["sources"][0]
    operations = []
    for identity, price in (("import-a", None), ("import-b", "0")):
        operations.extend(
            [
                dict(
                    action="device_put",
                    value=dict(
                        id=identity,
                        name="导入产品",
                        variant_id=variant["id"],
                        source_id=source["id"],
                        quantity="1.25",
                        kind="software",
                    ),
                ),
                dict(
                    action="purchase_set", device_id=identity, quantity="1.25", evidence="原表确认"
                ),
                dict(
                    action="price_set",
                    value=dict(
                        device_id=identity,
                        variant_id=variant["id"],
                        source_id=source["id"],
                        mode="import",
                        unit_price=price,
                        evidence="导入隔离原表",
                    ),
                ),
                dict(action="description_set", device_id=identity, text="项目说明"),
            ]
        )
    request = mutation(draft, operations=operations)
    updated = call(client, "list_update", request)
    assert call(client, "list_update", request) == updated
    quote = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    imported = [line for line in quote["items"] if line["device_id"].startswith("import-")]
    assert len(imported) == 2
    assert all(
        line["kind"] == "software" and line["purchase_quantity"] == "1.25" for line in imported
    )
    assert quote["total"] is None
    assert {line["unit_price"] for line in imported} == {None, "0"}
    call(
        client,
        "list_update",
        mutation(updated, operations=[dict(action="remove", collection="devices", id="import-a")]),
    )
    quote = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert "import-a" not in quote["metadata"]["descriptions"]
    assert quote["total"] == "30.86"
