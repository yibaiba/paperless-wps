from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook


def merged_book():
    book = Workbook()
    sheet = book.active
    sheet.title = "合并备注测试"
    sheet.append(["名称", "型号", "数量", "单价", "备注"])
    sheet.append(["输入设备", "INPUT", 2, "=1+2", "选配\n按现场确认"])
    sheet.append(["输出设备", "OUTPUT"])
    for region in ("C2:C3", "D2:D3", "E2:E3"):
        sheet.merge_cells(region)
    sheet.merge_cells("A5:E5")
    sheet["A5"] = "新增分区"
    other = book.create_sheet("普通表")
    other.append(["型号", "数量"])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_import_returns_merge_provenance_without_copying_quantities_or_formulas(client):
    result = client.post(
        "/api/quotation/import-cells", files={"file": ("merged.xlsx", merged_book())}
    )
    assert result.status_code == 200, result.text
    sheet, other = result.json()["sheets"]
    assert sheet["rows"][1] == ["输入设备", "INPUT", "2", "=1+2", "选配\n按现场确认"]
    assert sheet["rows"][2] == ["输出设备", "OUTPUT"]
    assert sheet["rows"][4] == ["新增分区"]
    merges = {m["range"]: m for m in sheet["merges"]}
    assert merges["E2:E3"] == dict(
        range="E2:E3", start_row=2, end_row=3, start_column=5, end_column=5
    )
    assert merges["A5:E5"]["end_column"] == 5
    assert other == dict(name="普通表", rows=[["型号", "数量"]])


def test_invalid_merge_reference_is_explicit_error(client):
    source, target = BytesIO(merged_book()), BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as output:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="E2:E3"', b'ref="invalid-region"')
            output.writestr(entry, content)
    response = client.post(
        "/api/quotation/import-cells", files={"file": ("invalid.xlsx", target.getvalue())}
    )
    assert response.status_code == 422
