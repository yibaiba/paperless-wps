from copy import deepcopy
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from presales.quotation.artifacts import FileArtifacts


def call(client, name, data):
    response = client.post("/api/list-tools/" + name, json=data)
    assert response.status_code == 200, response.text
    return response.json()


def mutation(draft, **extra):
    return dict(
        draft_id=draft["id"],
        expected_revision=draft["revision"],
        operation_id=str(uuid4()),
        **extra,
    )


def create(client):
    return call(
        client,
        "list_create",
        dict(
            name="隔离清单 MCP 验证",
            actor="测试 Agent",
            evidence="仅测试",
            operation_id=str(uuid4()),
        ),
    )


def device_operations(catalog, *, quantity="2.5", price="12.345"):
    variant, source = catalog["variants"][0], catalog["sources"][0]
    return [
        dict(action="room_put", value=dict(id="room", name="会议室")),
        dict(
            action="system_put",
            value=dict(id="paper", name="无纸化会议系统", kind="无纸化", room_id="room"),
        ),
        dict(
            action="requirement_put",
            value=dict(id="role", role="服务端", system_id="paper", device_id="server"),
        ),
        dict(
            action="device_put",
            value=dict(
                id="server",
                name="测试服务器",
                variant_id=variant["id"],
                source_id=source["id"],
                quantity=quantity,
                kind="hardware",
            ),
        ),
        dict(
            action="supply_set",
            device_id="server",
            allocations=[
                dict(
                    id="supply",
                    device_id="server",
                    quantity=quantity,
                    source="purchase",
                    evidence="测试采购",
                )
            ],
        ),
        dict(
            action="quotation_set",
            value=dict(
                project_name="隔离项目",
                customer="测试客户",
                price_column="甲方指导价",
                prices=[
                    dict(
                        device_id="server",
                        variant_id=variant["id"],
                        source_id=source["id"],
                        mode="manual",
                        unit_price=price,
                        evidence="测试人工价",
                    )
                ],
            ),
        ),
    ]


def saved_project(client, catalog):
    draft = create(client)
    draft = call(client, "list_update", mutation(draft, operations=device_operations(catalog)))
    checked = call(client, "list_check", mutation(draft))
    return call(
        client,
        "list_save",
        mutation(checked, expected_project_revision=0, fingerprint=checked["check_fingerprint"]),
    )


def test_edit_check_save_retry_and_export(client, catalog, tmp_path):
    client.app.state.artifact_files = FileArtifacts(tmp_path)
    draft = create(client)
    update = mutation(draft, operations=device_operations(catalog))
    edited = call(client, "list_update", update)
    assert call(client, "list_update", update) == edited
    quote = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert quote["total"] == "30.86"
    checked = call(client, "list_check", mutation(edited))
    save = mutation(checked, expected_project_revision=0, fingerprint=checked["check_fingerprint"])
    saved = call(client, "list_save", save)
    assert call(client, "list_save", save) == saved
    assert saved["project_revision"] == 1
    exported = call(
        client,
        "list_export",
        dict(project_id=saved["project_id"], revision=1, output="both", operation_id="export-1"),
    )
    assert len(exported["artifacts"]) == 2
    artifact = exported["artifacts"][1]
    response = client.get(artifact["download_path"])
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content), data_only=False)
    assert wb.active["H8"].value == '=IF(OR(E8="",G8=""),"",ROUND(E8*G8,2))'
    assert len(wb.active._images) == 1
    cached = load_workbook(BytesIO(response.content), data_only=True)
    assert cached.active["H35"].value == 30.86
    assert cached.active["C4"].value == "测试客户"


def test_missing_prices_stale_price_and_zero(client, catalog):
    draft = create(client)
    operations = device_operations(catalog, price="0")
    edited = call(client, "list_update", mutation(draft, operations=operations))
    assert call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))["total"] == "0.00"
    replacement = deepcopy(operations[3])
    replacement["value"].update(
        variant_id=catalog["variants"][1]["id"], source_id=catalog["sources"][1]["id"]
    )
    changed = call(client, "list_update", mutation(edited, operations=[replacement]))
    quote = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert quote["total"] is None
    assert "过期" in quote["items"][0]["issues"][0]
    call(
        client,
        "list_update",
        mutation(changed, operations=[dict(action="price_readopt", device_ids=["server"])]),
    )
    quote = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert "缺价" in quote["items"][0]["issues"][0]


def test_atomic_errors_versions_and_fixed_history(client, catalog):
    draft = create(client)
    bad = mutation(
        draft,
        operations=[
            dict(action="room_put", value=dict(id="room", name="不会保存")),
            dict(action="remove", collection="devices", id="missing"),
        ],
    )
    assert client.post("/api/list-tools/list_update", json=bad).status_code == 422
    current = call(client, "list_get", dict(draft_id=draft["id"]))
    assert current["rooms"] == []
    assert current["revision"] == draft["revision"]
    saved = saved_project(client, catalog)
    copied = call(
        client,
        "list_create",
        dict(
            name="改单",
            actor="测试",
            evidence="复制",
            operation_id="copy",
            project_id=saved["project_id"],
            revision=1,
        ),
    )
    copied = call(
        client,
        "list_update",
        mutation(copied, operations=[dict(action="remove", collection="devices", id="server")]),
    )
    checked = call(client, "list_check", mutation(copied))
    call(
        client,
        "list_save",
        mutation(checked, expected_project_revision=1, fingerprint=checked["check_fingerprint"]),
    )
    before = call(
        client, "list_get", dict(project_id=saved["project_id"], revision=1, view="devices")
    )
    after = call(
        client, "list_get", dict(project_id=saved["project_id"], revision=2, view="devices")
    )
    assert len(before["items"]) == 1 and after["items"] == []
    assert client.post("/api/list-tools/list_check", json=mutation(copied)).status_code == 409


def test_queries_and_no_default_quantity(client, catalog):
    draft = create(client)
    call(client, "systems_list", {})
    found = call(client, "catalog_search", dict(query="SERVER-X", limit=1))
    assert found["total"] == 2 and found["next_offset"] == 1
    detail = call(client, "catalog_get", dict(variant_id=found["items"][0]["variant_id"]))
    assert detail["sources"]
    operations = device_operations(catalog)
    del operations[3]["value"]["quantity"]
    assert (
        client.post(
            "/api/list-tools/list_update", json=mutation(draft, operations=operations)
        ).status_code
        == 422
    )
