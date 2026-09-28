from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from presales.quotation.artifacts import FileArtifacts
from presales.quotation.template import TEMPLATE_PATH, TEMPLATE_SHA256

from .conftest import AUTHOR
from .test_evolution_sharing import cross_system
from .test_list_mcp import call, create, device_operations, mutation


def test_template_grows_preserves_logo_and_formula_caches(client, catalog, tmp_path):
    client.app.state.artifact_files = FileArtifacts(tmp_path)
    operations = device_operations(catalog, quantity="1", price="1.235")
    for index in range(1, 31):
        identity = f"extra{index}"
        device = dict(operations[3]["value"], id=identity, name=f"扩行产品 {index}")
        operations.append(dict(action="device_put", value=device))
        operations.append(
            dict(
                action="supply_set",
                device_id=identity,
                allocations=[
                    dict(
                        id="supply-" + identity,
                        device_id=identity,
                        quantity="1",
                        source="purchase",
                        evidence="测试",
                    )
                ],
            )
        )
        quote = operations[5]["value"]
        quote["prices"].append(dict(quote["prices"][0], device_id=identity))
        quote.setdefault("sections", {})[identity] = (
            "无纸化会议系统"
            if index < 12
            else "会议扩声系统"
            if index < 23
            else "其他辅助设备"
            if index < 30
            else "会议预约"
        )
    operations[5]["value"]["customer"] = '=HYPERLINK("https://example.com")'
    draft = call(client, "list_update", mutation(create(client), operations=operations))
    checked = call(client, "list_check", mutation(draft))
    saved = call(
        client,
        "list_save",
        mutation(checked, expected_project_revision=0, fingerprint=checked["check_fingerprint"]),
    )
    artifact = call(
        client,
        "list_export",
        dict(
            project_id=saved["project_id"], revision=1, output="quotation", operation_id="expansion"
        ),
    )["artifacts"][0]
    content = client.get(artifact["download_path"]).content
    formulas = load_workbook(BytesIO(content), data_only=False).active
    values = load_workbook(BytesIO(content), data_only=True).active
    assert formulas.max_row > 35
    assert formulas["C4"].data_type == "s"
    assert values.cell(values.max_row, 8).value == 38.44
    assert (
        sum(1 for row in values if row[1].value and str(row[1].value).startswith("扩行产品")) == 30
    )
    assert any(row[0].value == "会议预约" for row in values)
    assert len(formulas._images) == 1
    with ZipFile(TEMPLATE_PATH) as original, ZipFile(BytesIO(content)) as exported:
        original_image = original.read("xl/media/image.jpg")
        media = [n for n in exported.namelist() if n.startswith("xl/media/")]
        assert exported.read(media[0]) == original_image
    assert sha256(TEMPLATE_PATH.read_bytes()).hexdigest() == TEMPLATE_SHA256


def test_shared_accessory_usage_existing_and_delete(client, catalog):
    config, _ = cross_system(client, catalog)
    project = client.post("/api/projects", json={"name": "配套共享隔离测试"}).json()
    client.put(
        "/api/configuration/projects/" + project["id"],
        json=dict(expected_revision=0, configuration=config),
    ).raise_for_status()
    draft = call(
        client,
        "list_create",
        dict(
            name="共享服务器", project_id=project["id"], revision=1, operation_id="shared", **AUTHOR
        ),
    )
    # Apply one server, then explicitly allocate the same instance to the second system.
    for index in range(2):
        draft = call(client, "list_check", mutation(draft))
        result = call(client, "list_get", dict(draft_id=draft["id"], view="issues"))
        suggestion = next(
            i for i in result["items"] if i.get("kind") == "accessory" and float(i["missing"]) > 0
        )
        operation = dict(
            action="accessory_apply",
            suggestion_id=suggestion["id"],
            fingerprint=draft["calculation_fingerprint"],
            quantity="1",
        )
        if index == 0:
            operation.update(
                variant_id=catalog["variants"][1]["id"],
                source_id=catalog["sources"][1]["id"],
                supply_source="existing",
                supply_evidence="测试客户已有",
            )
        else:
            devices = call(client, "list_get", dict(draft_id=draft["id"], view="devices"))["items"]
            operation["existing_device_id"] = next(
                d["device_id"] for d in devices if d["kind"] == "hardware"
            )
        draft = call(client, "list_update", mutation(draft, operations=[operation]))
    devices = call(client, "list_get", dict(draft_id=draft["id"], view="devices"))["items"]
    hardware = [d for d in devices if d["kind"] == "hardware"]
    assert len(hardware) == 1 and len(hardware[0]["consumers"]) == 2
    assert hardware[0]["supply"]["purchase"] == "0"
    issues = call(client, "list_get", dict(draft_id=draft["id"], view="issues"))["items"]
    assert any(i["kind"] == "sharing" and i["status"] == "unknown" for i in issues)
    call(
        client,
        "list_update",
        mutation(
            draft,
            operations=[dict(action="remove", collection="devices", id=hardware[0]["device_id"])],
        ),
    )
    issues = call(client, "list_get", dict(draft_id=draft["id"], view="issues"))["items"]
    assert len([i for i in issues if i["kind"] == "accessory" and i["missing"] == "1"]) == 2


def test_missing_price_export_and_fingerprint_invalidation(client, catalog, tmp_path):
    client.app.state.artifact_files = FileArtifacts(tmp_path)
    operations = device_operations(catalog)
    operations[5]["value"]["prices"] = []
    draft = call(client, "list_update", mutation(create(client), operations=operations))
    checked = call(client, "list_check", mutation(draft))
    changed = call(
        client,
        "list_update",
        mutation(
            checked, operations=[dict(action="room_put", value=dict(id="room", name="改房间"))]
        ),
    )
    stale = mutation(changed, expected_project_revision=0, fingerprint=checked["check_fingerprint"])
    assert client.post("/api/list-tools/list_save", json=stale).status_code == 409
    checked = call(client, "list_check", mutation(changed))
    saved = call(
        client,
        "list_save",
        mutation(checked, expected_project_revision=0, fingerprint=checked["check_fingerprint"]),
    )
    artifact = call(
        client,
        "list_export",
        dict(
            project_id=saved["project_id"],
            revision=1,
            operation_id="missing-price",
            output="quotation",
        ),
    )["artifacts"][0]
    ws = load_workbook(
        BytesIO(client.get(artifact["download_path"]).content), data_only=True
    ).active
    assert ws["H35"].value == "待确认" and ws["G8"].value is None
    assert "缺价" in ws["J8"].value
