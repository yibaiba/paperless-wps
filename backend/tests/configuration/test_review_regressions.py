from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from presales.quotation.calculation import with_quotation
from presales.quotation.excel import ExcelRenderer
from presales.quotation.template import TEMPLATE_PATH

from .test_list_mcp import call, create, device_operations, mutation, saved_project
from .test_web_drafts import write


def test_first_web_draft_can_check_and_save_unsaved_project(client, catalog):
    project = client.post("/api/projects", json={"name": "隔离首个网页版本"}).json()
    response = client.post(
        "/api/work-drafts",
        json=dict(project_id=project["id"], expected_revision=0, operation_id=str(uuid4())),
    )
    assert response.status_code == 200, response.text
    updated = write(client, response.json(), device_operations(catalog))
    assert updated.status_code == 200, updated.text
    checked = call(client, "list_check", mutation(updated.json()))
    request = mutation(
        checked, expected_project_revision=0, fingerprint=checked["check_fingerprint"]
    )
    saved = call(client, "list_save", request)
    assert saved["project_id"] == project["id"] and saved["project_revision"] == 1
    assert call(client, "list_save", request) == saved


def test_saved_unassigned_device_keeps_actionable_issue(client, catalog):
    operations = [op for op in device_operations(catalog) if op["action"] != "requirement_put"]
    draft = call(client, "list_update", mutation(create(client), operations=operations))
    checked = call(client, "list_check", mutation(draft))
    saved = call(
        client,
        "list_save",
        mutation(checked, expected_project_revision=0, fingerprint=checked["check_fingerprint"]),
    )
    reopened = client.get("/api/configuration/projects/" + saved["project_id"]).json()
    issues = [c for c in reopened["checks"] if c["kind"] == "assignment"]
    assert len(issues) == 1
    assert issues[0]["device_id"] == "server"
    assert issues[0]["action"]["type"] == "assign_device"
    assert reopened["readiness"]["ready_for_confirmation"] is False


def test_export_keeps_stale_description_pending_after_recalculation(client, catalog):
    saved = saved_project(client, catalog)
    record = client.get("/api/configuration/projects/" + saved["project_id"]).json()
    device = record["configuration"]["devices"][0]
    record["configuration"]["quotation"]["descriptions"][device["id"]] = dict(
        variant_id="previous-product", source_id=device["source_id"], text="旧型号说明"
    )
    checked = with_quotation(record)
    assert checked["quotation_output"]["total"] is None
    content = ExcelRenderer(TEMPLATE_PATH).quotation(checked)
    formulas = load_workbook(BytesIO(content)).active
    values = load_workbook(BytesIO(content), data_only=True).active
    # Price and quantity are valid: only the unresolved description prevents a total.
    assert isinstance(values["H8"].value, (int, float))
    assert ",TRUE)" in formulas.cell(formulas.max_row, 8).value
    assert values.cell(values.max_row, 8).value == "待确认"
