from io import BytesIO

from openpyxl import load_workbook

from presales.quotation.excel import ExcelRenderer
from presales.quotation.template import TEMPLATE_PATH
from presales.storage import ProductRecord

from .test_list_mcp import call, create, device_operations, mutation, saved_project


def test_long_descriptions_continue_without_extra_quantities(client, catalog):
    saved = saved_project(client, catalog)
    record = client.get("/api/configuration/projects/" + saved["project_id"]).json()
    description = "产品技术说明与功能参数，完整保留长文本。" * 200
    record["quotation_output"]["lines"][0]["specification"] = description
    content = ExcelRenderer(TEMPLATE_PATH).quotation(record)
    ws = load_workbook(BytesIO(content), data_only=True).active
    text = "".join(str(ws.cell(r, 4).value or "").replace("\n", "") for r in range(8, ws.max_row))
    assert text == description
    assert sum(ws.cell(r, 5).value or 0 for r in range(8, ws.max_row)) == 2.5
    assert sum(ws.cell(r, 8).value or 0 for r in range(8, ws.max_row)) == 30.86
    assert max(ws.row_dimensions[r].height for r in range(8, ws.max_row)) <= 360
    assert ws.cell(ws.max_row, 8).value == 30.86


def test_explicit_quote_group_and_unknown_supply(client, catalog):
    operations = device_operations(catalog)
    operations[4]["allocations"][0]["source"] = "unknown"
    operations[5]["value"]["sections"] = {"server": "独立无纸化分区"}
    draft = call(client, "list_update", mutation(create(client), operations=operations))
    result = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert result["total"] is None
    assert result["items"][0]["section"] == "独立无纸化分区"
    assert result["items"][0]["supply_complete"] is False
    assert result["known_subtotal"] == "0.00"
    operations[5]["value"]["sections"]["server"] = ""
    call(client, "list_update", mutation(draft, operations=[operations[5]]))
    result = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert result["items"][0]["section"] == "无纸化会议系统"


def test_saved_revision_search_and_validated_boundaries(client, catalog):
    saved = saved_project(client, catalog)
    found = call(client, "list_search", dict(query="MCP"))
    assert found["items"][0]["revision"] == 1
    history = call(client, "list_search", dict(project_id=saved["project_id"]))
    assert history["items"][0]["quotation_total"] == "30.86"
    read = call(client, "list_get", dict(project_id=saved["project_id"]))
    assert read["check_current"] is True
    for tool, arguments in [("catalog_get", {}), ("list_search", {"query": {}})]:
        assert client.post("/api/list-tools/" + tool, json=arguments).status_code == 422


def test_metadata_patch_retains_manual_prices(client, catalog):
    draft = call(
        client, "list_update", mutation(create(client), operations=device_operations(catalog))
    )
    call(
        client,
        "list_update",
        mutation(
            draft,
            operations=[
                dict(
                    action="quotation_set",
                    value=dict(customer="改单客户", price_column="出厂指导价"),
                )
            ],
        ),
    )
    result = call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert result["metadata"]["customer"] == "改单客户"
    assert result["items"][0]["price_selection"]["mode"] == "manual"
    assert result["total"] == "30.86"


def test_quote_patch_presence_is_part_of_idempotency_key(client, catalog):
    draft = call(
        client, "list_update", mutation(create(client), operations=device_operations(catalog))
    )
    update = mutation(draft, operations=[dict(action="quotation_set", value=dict(customer="客户"))])
    call(client, "list_update", update)
    update["operations"][0]["value"]["prices"] = []
    response = client.post("/api/list-tools/list_update", json=update)
    assert response.status_code == 409 and "IDEMPOTENCY_CONFLICT" in response.text


def test_price_columns_and_quantity_changes_are_explicit(client, catalog):
    with client.app.state.session_factory() as session:
        source = session.get(ProductRecord, catalog["sources"][0]["id"])
        source.payload = dict(
            source.payload,
            prices={"甲方指导价": "20.50", "市场参考报价": "18.555", "出厂指导价": "按项目申请"},
        )
        session.commit()
    operations = device_operations(catalog)
    operations[5]["value"]["prices"] = []
    draft = call(client, "list_update", mutation(create(client), operations=operations))
    def quote():
        return call(client, "list_get", dict(draft_id=draft["id"], view="quotation"))
    assert quote()["total"] == "51.25"
    draft = call(
        client,
        "list_update",
        mutation(
            draft,
            operations=[dict(action="quotation_set", value=dict(price_column="市场参考报价"))],
        ),
    )
    assert quote()["total"] == "46.39"
    operations[3]["value"]["quantity"] = "5"
    draft = call(client, "list_update", mutation(draft, operations=[operations[3]]))
    assert quote()["total"] is None and quote()["items"][0]["unknown_quantity"] == "2.5"
    operations[4]["allocations"][0]["quantity"] = "5"
    draft = call(client, "list_update", mutation(draft, operations=[operations[4]]))
    assert quote()["total"] == "92.78"
    call(
        client,
        "list_update",
        mutation(
            draft, operations=[dict(action="quotation_set", value=dict(price_column="出厂指导价"))]
        ),
    )
    assert quote()["total"] is None
